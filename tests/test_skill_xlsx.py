"""xlsx skill — load, eligibility, and create→inspect→edit round-trip."""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import pytest

from agentos.skills.eligibility import EligibilityContext, check_eligibility
from agentos.skills.loader import SkillLoader

ROOT = Path(__file__).resolve().parents[1]
BUNDLED = ROOT / "src" / "agentos" / "skills" / "bundled"
SCRIPTS = BUNDLED / "xlsx" / "scripts"


def _spec() -> object:
    return SkillLoader(bundled_dir=BUNDLED).get_by_name("xlsx")


def test_skill_loads() -> None:
    spec = _spec()
    assert spec is not None
    assert spec.name == "xlsx"
    assert spec.metadata is not None


def test_eligibility_with_python(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "agentos.skills.eligibility.shutil.which",
        lambda name: "/usr/bin/python3" if name in {"python", "python3"} else None,
    )
    spec = _spec()
    assert spec is not None
    assert check_eligibility(spec, EligibilityContext.auto())


@pytest.mark.parametrize(
    "merged",
    [
        pytest.param([{"range": "A1:B1"}], id="dictionary"),
        pytest.param(["A1:B1"], id="string"),
    ],
)
def test_round_trip_with_formula_and_merge(tmp_path: Path, merged: list[object]) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    spec = {
        "sheets": [
            {
                "name": "Sales",
                "rows": [
                    ["Region", "Revenue"],
                    ["NA", 1_200_000],
                    ["EU", 850_000],
                    ["Total", "=SUM(B2:B3)"],
                ],
                "merged": merged,
                "freeze": "A2",
            }
        ]
    }
    src = tmp_path / "book.xlsx"
    create_xlsx.build(spec).save(str(src))
    assert src.exists()

    inspected = inspect_xlsx.inspect(src, data_only=False)
    sheet = next(s for s in inspected["sheets"] if s["name"] == "Sales")
    assert sheet["max_row"] == 4
    assert sheet["max_col"] == 2
    assert sheet["merged"] == ["A1:B1"]
    assert sheet["freeze"] == "A2"

    last_row = sheet["rows"][3]
    assert last_row[1]["type"] == "f"
    assert last_row[1]["value"] == "=SUM(B2:B3)"

    from openpyxl import load_workbook

    wb = load_workbook(str(src))
    edit_xlsx.apply_ops(
        wb,
        [
            {"op": "set_cell", "sheet": "Sales", "row": 2, "col": 1, "value": "Americas"},
            {"op": "rename_sheet", "old": "Sales", "new": "Q3"},
        ],
    )
    out = tmp_path / "out.xlsx"
    wb.save(str(out))

    re_inspected = inspect_xlsx.inspect(out, data_only=False)
    sheet_q3 = next(s for s in re_inspected["sheets"] if s["name"] == "Q3")
    assert sheet_q3["rows"][1][0]["value"] == "Americas"


@pytest.mark.parametrize(
    ("merged", "expected"),
    [
        pytest.param(["A1:B1", {"range": "A3:B3"}], ["A1:B1", "A3:B3"], id="mixed"),
        pytest.param([], [], id="empty"),
        pytest.param(None, [], id="null"),
        pytest.param([None, 42, {}, {"other": "A1:B1"}], [], id="unsupported-entries"),
    ],
)
def test_create_merge_formats(tmp_path: Path, merged: object, expected: list[str]) -> None:
    from agentos.skills.bundled.xlsx.scripts.create_xlsx import build
    from agentos.skills.bundled.xlsx.scripts.inspect_xlsx import inspect

    path = tmp_path / "merges.xlsx"
    wb = build({"sheets": [{"name": "Merges", "merged": merged}, {"name": "Plain"}]})
    wb.save(path)
    wb.close()

    sheets = inspect(path, data_only=False)["sheets"]
    assert sheets[0]["merged"] == expected
    assert sheets[1]["merged"] == []


@pytest.mark.parametrize("merged", [["not-a-range"], [{"range": "not-a-range"}]])
def test_create_rejects_invalid_merge_ranges(merged: list[object]) -> None:
    from agentos.skills.bundled.xlsx.scripts.create_xlsx import build

    with pytest.raises(ValueError, match="not a valid coordinate or range"):
        build({"sheets": [{"merged": merged}]})


def test_create_accepts_inspected_merge_ranges(tmp_path: Path) -> None:
    from openpyxl import Workbook

    from agentos.skills.bundled.xlsx.scripts.create_xlsx import build
    from agentos.skills.bundled.xlsx.scripts.inspect_xlsx import inspect

    original = Workbook()
    ws = original.active
    assert ws is not None
    ws.merge_cells("A1:C1")
    ws.merge_cells("B3:B5")
    source = tmp_path / "source.xlsx"
    original.save(source)
    original.close()
    merges = inspect(source, data_only=False)["sheets"][0]["merged"]
    assert merges == ["A1:C1", "B3:B5"]

    # Reuse only merge metadata: inspector rows have a different schema.
    rebuilt = build({"sheets": [{"merged": merges}]})
    target = tmp_path / "rebuilt.xlsx"
    rebuilt.save(target)
    rebuilt.close()
    assert inspect(target, data_only=False)["sheets"][0]["merged"] == merges


def test_text_escapes_formula(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["a"]]}]}).save(str(src))

    from openpyxl import load_workbook

    wb = load_workbook(str(src))
    edit_xlsx.apply_ops(
        wb,
        [
            {
                "op": "set_cell",
                "sheet": "S",
                "row": 2,
                "col": 1,
                "value": "=hello",
                "as_text": True,
            },
        ],
    )
    out = tmp_path / "out.xlsx"
    wb.save(str(out))

    inspected = inspect_xlsx.inspect(out, data_only=False)
    sheet = inspected["sheets"][0]
    cell_value = sheet["rows"][1][0]["value"]
    assert isinstance(cell_value, str)
    assert cell_value.lstrip("'") == "=hello"


def test_inspect_xlsx_creates_parent_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["a"]]}]}).save(str(src))

    out = tmp_path / "nested" / "dir" / "out.json"
    monkeypatch.setattr(sys, "argv", ["inspect_xlsx.py", str(src), "--out", str(out)])
    assert inspect_xlsx.main() == 0
    assert out.is_file()


def _edit_and_reload(tmp_path: Path, ops: list[dict[str, object]]) -> object:
    """Run ``ops`` against a one-cell workbook and reload the saved result."""
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    from openpyxl import load_workbook

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["a"]]}]}).save(str(src))
    wb = load_workbook(str(src))
    edit_xlsx.apply_ops(wb, ops)
    out = tmp_path / "out.xlsx"
    wb.save(str(out))
    return load_workbook(str(out))["S"]


def _set_cell(row: int, value: object, **extra: object) -> dict[str, object]:
    return {"op": "set_cell", "sheet": "S", "row": row, "col": 1, "value": value, **extra}


def test_as_text_keeps_the_formula_string_out_of_the_cell_value(tmp_path: Path) -> None:
    """``as_text: true`` stores ``=hello``, not Excel's input-mode apostrophe.

    The apostrophe is an escape typed into Excel's formula bar, not cell
    content. Prepending it left the cell holding ``'=hello`` — seven characters
    where the caller passed six — so a round-trip comparison against the
    original string failed and both Excel and LibreOffice rendered a stray
    apostrophe.
    """
    sheet = _edit_and_reload(tmp_path, [_set_cell(2, "=hello", as_text=True)])
    cell = sheet.cell(row=2, column=1)

    assert cell.value == "=hello"
    assert cell.data_type == "s"
    assert cell.quotePrefix is True


def test_inspect_reports_the_escaped_formula_without_an_apostrophe(
    tmp_path: Path,
) -> None:
    """The skill's own inspector is where a caller reads the value back.

    ``inspect_xlsx`` reported ``'=hello``, so the corruption was visible
    through the documented read path and not only through openpyxl directly.
    """
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    from openpyxl import load_workbook

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["a"]]}]}).save(str(src))
    wb = load_workbook(str(src))
    edit_xlsx.apply_ops(wb, [_set_cell(2, "=hello", as_text=True)])
    out = tmp_path / "out.xlsx"
    wb.save(str(out))

    reported = inspect_xlsx.inspect(out, data_only=False)["sheets"][0]["rows"][1][0]

    assert reported["value"] == "=hello"
    assert reported["type"] == "s"


def test_without_as_text_a_formula_string_stays_a_formula(tmp_path: Path) -> None:
    """The default path is unchanged: openpyxl still records a formula cell."""
    sheet = _edit_and_reload(tmp_path, [_set_cell(2, "=SUM(A1:A1)")])
    cell = sheet.cell(row=2, column=1)

    assert cell.value == "=SUM(A1:A1)"
    assert cell.data_type == "f"
    assert cell.quotePrefix is False


def test_as_text_suppresses_the_iso_datetime_coercion(tmp_path: Path) -> None:
    """An explicit ``as_text`` request must win over the ISO-8601 heuristic.

    ``_coerce`` only consulted the flag on the formula branch, so there was no
    way to store a timestamp as text: every 19-character string with ``T`` at
    index 10 became a ``datetime`` cell regardless.
    """
    sheet = _edit_and_reload(tmp_path, [_set_cell(2, "2026-05-06T09:00:00", as_text=True)])
    cell = sheet.cell(row=2, column=1)

    assert cell.value == "2026-05-06T09:00:00"
    assert isinstance(cell.value, str)
    assert cell.data_type == "s"


def test_without_as_text_an_iso_string_still_becomes_a_datetime(tmp_path: Path) -> None:
    """The documented datetime convenience keeps working untouched."""
    from datetime import datetime

    sheet = _edit_and_reload(tmp_path, [_set_cell(2, "2026-05-06T09:00:00")])
    cell = sheet.cell(row=2, column=1)

    assert cell.value == datetime(2026, 5, 6, 9, 0)
    assert cell.data_type == "d"


def test_as_text_does_not_flag_an_ordinary_string(tmp_path: Path) -> None:
    """``quotePrefix`` only stands in for the escape a formula string needs.

    A plain string must not pick up the flag, or every text cell written with
    ``as_text`` would carry a style Excel attributes to manual escaping.
    """
    sheet = _edit_and_reload(tmp_path, [_set_cell(2, "plain", as_text=True)])
    cell = sheet.cell(row=2, column=1)

    assert cell.value == "plain"
    assert cell.data_type == "s"
    assert cell.quotePrefix is False


def test_as_text_leaves_non_string_values_alone(tmp_path: Path) -> None:
    """Numbers and booleans keep their own cell types; the flag is about text."""
    sheet = _edit_and_reload(
        tmp_path, [_set_cell(2, 42, as_text=True), _set_cell(3, True, as_text=True)]
    )

    assert sheet.cell(row=2, column=1).value == 42
    assert sheet.cell(row=2, column=1).data_type == "n"
    assert sheet.cell(row=3, column=1).value is True
    assert sheet.cell(row=3, column=1).data_type == "b"


def test_the_apostrophe_escape_and_as_text_produce_the_same_cell(tmp_path: Path) -> None:
    """``SKILL.md`` offers two spellings of one request; they must agree.

    Excel's leading apostrophe is the input escape for a formula-looking
    value, so ``{"value": "'=hello", "as_text": true}`` asks for exactly what
    ``{"value": "=hello", "as_text": true}`` asks for. Storing the apostrophe
    as data left the two spellings on different cells -- one holding ``=hello``
    with ``quotePrefix`` set, the other holding a literal ``'=hello``.
    """
    sheet = _edit_and_reload(
        tmp_path,
        [_set_cell(2, "=hello", as_text=True), _set_cell(3, "'=hello", as_text=True)],
    )
    plain = sheet.cell(row=2, column=1)
    escaped = sheet.cell(row=3, column=1)

    assert (escaped.value, escaped.data_type, escaped.quotePrefix) == (
        plain.value,
        plain.data_type,
        plain.quotePrefix,
    )
    assert escaped.value == "=hello"
    assert escaped.quotePrefix is True


def test_a_value_that_genuinely_starts_with_an_apostrophe_keeps_it(tmp_path: Path) -> None:
    """The escape is only consumed when it escapes a formula.

    Stripping every leading apostrophe would turn ``'tis`` into ``tis`` --
    trading the reported bug for a quieter one.
    """
    sheet = _edit_and_reload(tmp_path, [_set_cell(2, "'tis", as_text=True)])
    cell = sheet.cell(row=2, column=1)

    assert cell.value == "'tis"
    assert cell.data_type == "s"
    assert cell.quotePrefix is False


def test_the_apostrophe_escape_is_not_consumed_without_as_text(tmp_path: Path) -> None:
    """Without the flag nothing is interpreted; the value is stored verbatim."""
    sheet = _edit_and_reload(tmp_path, [_set_cell(2, "'=hello")])

    assert sheet.cell(row=2, column=1).value == "'=hello"


def _import_scripts() -> tuple[Any, Any, Any]:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return create_xlsx, edit_xlsx, inspect_xlsx


def _run_cli(
    edit_xlsx: Any,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    src: Path,
    out: Path,
    ops: list[dict[str, Any]],
    tmp_path: Path,
) -> dict[str, Any]:
    """Drive the script the way a caller does: ops file in, workbook out.

    In-process rather than a subprocess, matching how the other bundled-skill
    tests are run on Windows. Going through ``main`` is what matters here: the
    reported ``applied`` count is part of the contract being fixed, and reading
    the saved file back is the only way to see whether a write really happened
    -- the in-memory workbook would look right either way.
    """
    ops_path = tmp_path / "ops.json"
    ops_path.write_text(json.dumps(ops), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", ["edit_xlsx.py", str(src), str(ops_path), "--out", str(out)])
    assert edit_xlsx.main() == 0
    return dict(json.loads(capsys.readouterr().out.strip()))


def _cell(path: Path, row: int, col: int, sheet: str = "S") -> Any:
    """Read one cell from the saved file.

    Addressed directly rather than through ``inspect``: clearing the only
    populated cell leaves the sheet with no used range at all, so a row-indexed
    read would raise IndexError instead of reporting the value as empty.
    """
    from openpyxl import load_workbook

    return load_workbook(str(path))[sheet].cell(row=row, column=col).value


def test_set_cell_explicit_null_clears_the_cell(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    create_xlsx, edit_xlsx, inspect_xlsx = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["keep me"]]}]}).save(str(src))

    out = tmp_path / "out.xlsx"
    report = _run_cli(
        edit_xlsx,
        monkeypatch,
        capsys,
        src,
        out,
        [{"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": None}],
        tmp_path,
    )

    assert report == {"applied": 1}
    assert _cell(out, 1, 1) is None


def test_set_cell_without_a_value_key_is_skipped(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    # A malformed op must not read as "clear this cell", and must not be
    # counted -- otherwise a typo silently wipes data and reports success.
    create_xlsx, edit_xlsx, inspect_xlsx = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["keep me"]]}]}).save(str(src))

    out = tmp_path / "out.xlsx"
    report = _run_cli(
        edit_xlsx,
        monkeypatch,
        capsys,
        src,
        out,
        [{"op": "set_cell", "sheet": "S", "row": 1, "col": 1}],
        tmp_path,
    )

    assert report == {"applied": 0}
    assert _cell(out, 1, 1) == "keep me"


@pytest.mark.parametrize(("value", "expected"), [(0, 0), (False, False)])
def test_set_cell_writes_falsy_values(
    value: Any,
    expected: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    # Guards the fix: these are values, not absence. The type assertion is the
    # point -- False must stay a boolean rather than collapse to 0.
    create_xlsx, edit_xlsx, _ = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["old"]]}]}).save(str(src))

    out = tmp_path / "out.xlsx"
    report = _run_cli(
        edit_xlsx,
        monkeypatch,
        capsys,
        src,
        out,
        [{"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": value}],
        tmp_path,
    )

    assert report == {"applied": 1}
    written = _cell(out, 1, 1)
    assert written == expected
    assert isinstance(written, type(expected))


def test_set_cell_empty_string_counts_as_an_edit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    # Excel has no empty-string cell, so openpyxl round-trips "" as None. That
    # is unchanged by this fix and is why "" is not a substitute for an
    # explicit null: the op still counts as applied either way, but only the
    # null path is documented as clearing.
    create_xlsx, edit_xlsx, _ = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["old"]]}]}).save(str(src))

    out = tmp_path / "out.xlsx"
    report = _run_cli(
        edit_xlsx,
        monkeypatch,
        capsys,
        src,
        out,
        [{"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": ""}],
        tmp_path,
    )

    assert report == {"applied": 1}
    assert _cell(out, 1, 1) is None


def test_clearing_a_cell_keeps_its_style(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    from openpyxl import load_workbook

    create_xlsx, edit_xlsx, _ = _import_scripts()
    src = tmp_path / "book.xlsx"
    wb = create_xlsx.build({"sheets": [{"name": "S", "rows": [["styled"]]}]})
    wb["S"].cell(row=1, column=1).number_format = "0.00%"
    wb.save(str(src))

    out = tmp_path / "out.xlsx"
    _run_cli(
        edit_xlsx,
        monkeypatch,
        capsys,
        src,
        out,
        [{"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": None}],
        tmp_path,
    )

    cell = load_workbook(str(out))["S"].cell(row=1, column=1)
    assert cell.value is None
    assert cell.number_format == "0.00%"


@pytest.mark.parametrize("spec", [None, [], "not-a-dict", 42, ("also", "not-a-dict")])
def test_build_returns_a_base_workbook_for_a_non_dict_spec(spec: object) -> None:
    """A non-dict spec must not crash `.get()` -- it never had a `sheets` key to read."""
    create_xlsx, _, _ = _import_scripts()

    wb = create_xlsx.build(spec)

    assert wb.sheetnames == ["Sheet"]


@pytest.mark.parametrize(
    "raw_sheets",
    [
        pytest.param("not-a-list", id="string"),
        pytest.param(None, id="null"),
        pytest.param([1, "x", None], id="list-of-non-dicts"),
    ],
)
def test_build_returns_a_base_workbook_for_malformed_sheets(raw_sheets: object) -> None:
    create_xlsx, _, _ = _import_scripts()

    wb = create_xlsx.build({"sheets": raw_sheets})

    assert wb.sheetnames == ["Sheet"]


def test_build_accepts_a_tuple_of_sheets_and_rows() -> None:
    """`sheets`/`rows` are consumed as any sequence, not specifically `list`."""
    create_xlsx, _, _ = _import_scripts()

    wb = create_xlsx.build({"sheets": ({"name": "S", "rows": (["a", "b"], ["c", "d"])},)})

    ws = wb["S"]
    assert ws.max_row == 2
    assert [ws.cell(row=r, column=c).value for r in (1, 2) for c in (1, 2)] == [
        "a",
        "b",
        "c",
        "d",
    ]


def test_build_normalizes_scalar_rows_into_single_cell_rows() -> None:
    """A scalar row becomes one cell, not an iteration over its characters/digits."""
    create_xlsx, _, _ = _import_scripts()

    wb = create_xlsx.build({"sheets": [{"name": "S", "rows": ["header", 42, None, ["a", "b"]]}]})

    ws = wb["S"]
    assert ws.max_row == 4
    assert ws.cell(row=1, column=1).value == "header"
    assert ws.cell(row=1, column=2).value is None  # not fragmented into 'h','e','a',...
    assert ws.cell(row=2, column=1).value == 42
    assert ws.cell(row=3, column=1).value is None
    assert ws.cell(row=4, column=1).value == "a"
    assert ws.cell(row=4, column=2).value == "b"


def test_build_skips_merge_entries_with_a_non_string_range() -> None:
    """A `range` of the wrong type is skipped rather than reaching `merge_cells`."""
    create_xlsx, _, _ = _import_scripts()

    wb = create_xlsx.build({"sheets": [{"merged": [{"range": 123}, {"range": ["A1:B1"]}]}]})

    assert list(wb["Sheet1"].merged_cells.ranges) == []


def test_create_xlsx_cli_reports_invalid_json_with_exit_code_2(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    create_xlsx, _, _ = _import_scripts()

    bad_json = tmp_path / "bad.json"
    bad_json.write_text("{not valid json", encoding="utf-8")
    out = tmp_path / "out.xlsx"
    monkeypatch.setattr(sys, "argv", ["create_xlsx.py", str(bad_json), "--out", str(out)])

    assert create_xlsx.main() == 2
    assert "invalid JSON spec" in capsys.readouterr().err
    assert not out.exists()


def test_create_xlsx_cli_reports_non_object_json_with_exit_code_2(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    create_xlsx, _, _ = _import_scripts()

    array_spec = tmp_path / "spec.json"
    array_spec.write_text("[1, 2, 3]", encoding="utf-8")
    out = tmp_path / "out.xlsx"
    monkeypatch.setattr(sys, "argv", ["create_xlsx.py", str(array_spec), "--out", str(out)])

    assert create_xlsx.main() == 2
    assert "JSON spec must be an object" in capsys.readouterr().err
    assert not out.exists()


# ── a malformed "sheets" shape is reported, not silently dropped (#silent-empty-xlsx) ──
#
# `build()` filters out anything that is not a well-formed sheet object --
# that permissiveness is deliberate and is exercised directly by
# `test_build_returns_a_base_workbook_for_malformed_sheets` above, since other
# callers use `build()` without going through the CLI's validation gate. The
# CLI entrypoint is a different question: `{"sheets": [1, 2, 3]}` and
# `{"sheets": "Sales"}` used to reach `build()` unchecked, which filtered every
# entry away and silently wrote a valid-looking single-default-sheet workbook
# at exit 0 -- indistinguishable from an intentional blank spec, and with the
# caller's actual sheets discarded without a trace. `main()` now validates the
# shape before calling `build()`, mirroring `check_body_entries` in the
# sibling `create_docx.py` script.


@pytest.mark.parametrize(
    ("raw_sheets", "expected_message"),
    [
        pytest.param([1, 2, 3], "sheet entry 0 must be an object, got int", id="all-non-dict"),
        pytest.param("Sales", '"sheets" must be an array, got str', id="not-a-list"),
        pytest.param(
            [{"name": "Good", "rows": [["a"]]}, "garbage"],
            "sheet entry 1 must be an object, got str",
            id="partially-malformed",
        ),
    ],
)
def test_create_xlsx_cli_rejects_a_malformed_sheets_shape(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    raw_sheets: object,
    expected_message: str,
) -> None:
    create_xlsx, _, _ = _import_scripts()

    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps({"sheets": raw_sheets}), encoding="utf-8")
    out = tmp_path / "out.xlsx"
    monkeypatch.setattr(sys, "argv", ["create_xlsx.py", str(spec_path), "--out", str(out)])

    assert create_xlsx.main() == 2

    captured = capsys.readouterr()
    assert captured.err.startswith("error: ")
    assert expected_message in captured.err
    assert not out.exists()


def test_create_xlsx_cli_a_missing_sheets_key_is_still_a_valid_blank_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """An omitted "sheets" key is a deliberate blank workbook, not malformed input."""
    create_xlsx, _, _ = _import_scripts()

    spec_path = tmp_path / "spec.json"
    spec_path.write_text("{}", encoding="utf-8")
    out = tmp_path / "out.xlsx"
    monkeypatch.setattr(sys, "argv", ["create_xlsx.py", str(spec_path), "--out", str(out)])

    assert create_xlsx.main() == 0
    assert out.is_file()


def test_create_xlsx_cli_still_creates_a_valid_spec(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    create_xlsx, _, inspect_xlsx = _import_scripts()

    spec_path = tmp_path / "spec.json"
    spec_path.write_text(
        json.dumps({"sheets": [{"name": "Sales", "rows": [["Region", "Revenue"], ["NA", 100]]}]}),
        encoding="utf-8",
    )
    out = tmp_path / "out.xlsx"
    monkeypatch.setattr(sys, "argv", ["create_xlsx.py", str(spec_path), "--out", str(out)])

    assert create_xlsx.main() == 0
    sheets = inspect_xlsx.inspect(out, data_only=False)["sheets"]
    assert sheets[0]["name"] == "Sales"
    assert sheets[0]["rows"][1][1]["value"] == 100


# ── unusable ops files are reported, not swallowed (#2145, #2146) ────────────


@pytest.mark.parametrize(
    ("body", "expected"),
    [
        ("not json at all", "is not valid JSON"),
        (
            '{"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": "x"}',
            "must be a JSON array",
        ),
        ('["set_cell"]', "op 0 must be an object"),
        (
            '[{"op": "set-cell", "sheet": "S", "row": 1, "col": 1, "value": "x"}]',
            "unknown kind 'set-cell'",
        ),
        ('[{"sheet": "S", "row": 1, "col": 1, "value": "x"}]', "unknown kind None"),
        (
            '[{"op": "rename_sheet", "old": "S", "new": "T"}, {"op": "merge-cells"}]',
            "op 1 has unknown kind 'merge-cells'",
        ),
    ],
)
def test_edit_xlsx_reports_an_unusable_ops_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    body: str,
    expected: str,
) -> None:
    """`json.loads` was unguarded and `ops = raw if isinstance(raw, list) else []`
    turned a wrong-shaped file into zero ops, writing an unchanged copy at exit 0."""
    create_xlsx, edit_xlsx, _inspect = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["keep me"]]}]}).save(str(src))
    ops = tmp_path / "ops.json"
    ops.write_text(body, encoding="utf-8")
    out = tmp_path / "out.xlsx"
    monkeypatch.setattr(sys, "argv", ["edit_xlsx.py", str(src), str(ops), "--out", str(out)])

    assert edit_xlsx.main() == 2

    captured = capsys.readouterr()
    assert captured.err.startswith("error: ")
    assert "Traceback" not in captured.err
    assert expected in captured.err
    if "unknown kind" in expected:
        assert "expected one of set_cell, rename_sheet, merge_cells" in captured.err
    assert captured.out == ""
    assert not out.exists()


def test_edit_xlsx_reports_an_ops_file_that_is_not_utf8(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    create_xlsx, edit_xlsx, _inspect = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["x"]]}]}).save(str(src))
    ops = tmp_path / "ops.json"
    ops.write_bytes(b"\xff\xfe[]")
    out = tmp_path / "out.xlsx"
    monkeypatch.setattr(sys, "argv", ["edit_xlsx.py", str(src), str(ops), "--out", str(out)])

    assert edit_xlsx.main() == 2
    assert "is not valid JSON" in capsys.readouterr().err
    assert not out.exists()


def test_edit_xlsx_leaves_an_existing_output_alone_when_the_ops_are_unusable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    create_xlsx, edit_xlsx, _inspect = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["source"]]}]}).save(str(src))
    out = tmp_path / "out.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["PREVIOUS WORK"]]}]}).save(str(out))
    before = out.read_bytes()
    ops = tmp_path / "ops.json"
    ops.write_text('[{"op": "set-cell", "sheet": "S", "row": 1, "col": 1}]', encoding="utf-8")
    monkeypatch.setattr(sys, "argv", ["edit_xlsx.py", str(src), str(ops), "--out", str(out)])

    assert edit_xlsx.main() == 2
    assert out.read_bytes() == before


def test_edit_xlsx_validates_the_ops_before_opening_the_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A bad ops file is named even when the input is not a workbook openpyxl
    can open: the ops are checked first."""
    _create, edit_xlsx, _inspect = _import_scripts()
    src = tmp_path / "book.xlsx"
    src.write_bytes(b"this is not a workbook")
    ops = tmp_path / "ops.json"
    ops.write_text('[{"op": "set-cell"}]', encoding="utf-8")
    monkeypatch.setattr(
        sys, "argv", ["edit_xlsx.py", str(src), str(ops), "--out", str(tmp_path / "o.xlsx")]
    )

    assert edit_xlsx.main() == 2


def test_edit_xlsx_load_ops_returns_the_ops_unchanged(tmp_path: Path) -> None:
    _create, edit_xlsx, _inspect = _import_scripts()
    ops = [{"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": "x", "as_text": True}]
    path = tmp_path / "ops.json"
    path.write_text(json.dumps(ops), encoding="utf-8")

    assert edit_xlsx.load_ops(path) == ops


def test_edit_xlsx_op_kinds_match_what_apply_ops_handles(tmp_path: Path) -> None:
    """Every advertised kind is applied; the allow-list and the dispatcher
    cannot drift apart without this failing."""
    create_xlsx, edit_xlsx, _inspect = _import_scripts()
    wb = create_xlsx.build({"sheets": [{"name": "S", "rows": [["a", "b"]]}]})
    ops = [
        {"op": "set_cell", "sheet": "S", "row": 2, "col": 1, "value": "c"},
        {"op": "merge_cells", "sheet": "S", "range": "A1:B1"},
        {"op": "rename_sheet", "old": "S", "new": "T"},
    ]
    assert {op["op"] for op in ops} == set(edit_xlsx.OP_KINDS)

    assert edit_xlsx.apply_ops(wb, ops) == 3


def test_edit_xlsx_an_empty_ops_array_is_still_accepted(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """``[]`` is a valid list of zero ops: the workbook is copied, applied is 0."""
    create_xlsx, edit_xlsx, _inspect = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["x"]]}]}).save(str(src))
    out = tmp_path / "out.xlsx"

    report = _run_cli(edit_xlsx, monkeypatch, capsys, src, out, [], tmp_path)

    assert report == {"applied": 0}
    assert _cell(out, 1, 1) == "x"


def test_edit_xlsx_a_skipped_valid_op_still_writes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """Validation is about the file's shape, not whether an op lands: a known
    op that names a missing sheet is still skipped, not rejected."""
    create_xlsx, edit_xlsx, _inspect = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["x"]]}]}).save(str(src))
    out = tmp_path / "out.xlsx"

    report = _run_cli(
        edit_xlsx,
        monkeypatch,
        capsys,
        src,
        out,
        [{"op": "set_cell", "sheet": "NOPE", "row": 1, "col": 1, "value": "y"}],
        tmp_path,
    )

    assert report == {"applied": 0}
    assert out.is_file()


def test_edit_xlsx_still_applies_a_valid_ops_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    create_xlsx, edit_xlsx, _inspect = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["before"]]}]}).save(str(src))
    out = tmp_path / "out.xlsx"
    report = _run_cli(
        edit_xlsx,
        monkeypatch,
        capsys,
        src,
        out,
        [{"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": "after"}],
        tmp_path,
    )

    assert report == {"applied": 1}
    assert _cell(out, 1, 1) == "after"

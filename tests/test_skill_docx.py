"""docx skill — load, eligibility, and create→inspect round-trip."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

from agentos.skills.eligibility import EligibilityContext, check_eligibility
from agentos.skills.loader import SkillLoader

ROOT = Path(__file__).resolve().parents[1]
BUNDLED = ROOT / "src" / "agentos" / "skills" / "bundled"
DOCX_DIR = BUNDLED / "docx"
SCRIPTS = DOCX_DIR / "scripts"


def _spec_to_loader() -> object:
    return SkillLoader(bundled_dir=BUNDLED).get_by_name("docx")


def test_skill_loads() -> None:
    spec = _spec_to_loader()
    assert spec is not None
    assert spec.name == "docx"
    assert spec.metadata is not None
    assert spec.provenance.origin == "clawhub-mit0"
    assert spec.provenance.license == "MIT-0"


def test_eligibility_with_python_present(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "agentos.skills.eligibility.shutil.which",
        lambda name: "/usr/bin/python3" if name in {"python", "python3"} else None,
    )
    spec = _spec_to_loader()
    assert spec is not None
    assert check_eligibility(spec, EligibilityContext.auto())


def test_eligibility_without_python(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "agentos.skills.eligibility.shutil.which",
        lambda name: None,
    )
    spec = _spec_to_loader()
    assert spec is not None
    assert not check_eligibility(spec, EligibilityContext.auto())


def test_create_then_inspect_round_trip(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    spec = {
        "metadata": {"title": "Round-trip", "author": "Tester"},
        "body": [
            {"kind": "heading", "level": 1, "text": "Hello"},
            {"kind": "paragraph", "text": "World."},
            {"kind": "table", "rows": [["A", "B"], ["1", "2"]]},
        ],
    }
    out_path = tmp_path / "out.docx"
    doc = create_docx.build(spec)
    doc.save(str(out_path))
    assert out_path.exists()

    inspected = inspect_docx.inspect(out_path)
    assert inspected["sections"] >= 1
    texts = [p["text"] for p in inspected["paragraphs"]]
    assert "Hello" in texts
    assert "World." in texts
    assert inspected["tables"] and inspected["tables"][0][0] == ["A", "B"]
    assert inspected["has_tracked_changes"] is False


def test_edit_replace_text(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import edit_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "src.docx"
    create_docx.build({"body": [{"kind": "paragraph", "text": "Hello {{NAME}}, welcome."}]}).save(
        str(src)
    )

    from docx import Document

    doc = Document(str(src))
    ops = [{"op": "replace_text", "find": "{{NAME}}", "with": "Wei"}]
    edit_docx.apply_ops(doc, ops)
    out = tmp_path / "out.docx"
    doc.save(str(out))

    inspected = inspect_docx.inspect(out)
    text = " ".join(p["text"] for p in inspected["paragraphs"])
    assert "{{NAME}}" not in text
    assert "Wei" in text


def test_inspect_cli_outputs_json(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "src.docx"
    create_docx.build({"body": [{"kind": "paragraph", "text": "x"}]}).save(str(src))

    payload = inspect_docx.inspect(src)
    encoded = json.dumps(payload, ensure_ascii=False)
    assert "paragraphs" in encoded
    assert "tables" in encoded


def _inspect_docx_module() -> object:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return inspect_docx


def test_inspect_survives_irregular_vertical_merge(tmp_path: Path) -> None:
    """`row.cells` resolves a `vMerge=continue` cell against the row above and
    raises `ValueError` when no cell starts at that grid offset there -- a
    layout non-Word generators produce. Walking the `<w:tc>` elements directly
    never enters that path. (#2154)"""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    inspect_docx = _inspect_docx_module()
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "Cell"
    continue_marker = OxmlElement("w:vMerge")
    continue_marker.set(qn("w:val"), "continue")
    table.cell(1, 1)._tc.get_or_add_tcPr().append(continue_marker)
    src = tmp_path / "irregular.docx"
    doc.save(str(src))

    tables = inspect_docx.inspect(src)["tables"]

    assert len(tables[0]) == 2
    assert tables[0][1][0] == "Cell"


def test_inspect_visits_a_merged_cell_once(tmp_path: Path) -> None:
    """`row.cells` repeats a horizontally merged cell once per grid column it
    spans, triplicating its text in the inspect output. (#2154)"""
    from docx import Document

    inspect_docx = _inspect_docx_module()
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "HELLO"
    src = tmp_path / "merged.docx"
    doc.save(str(src))

    tables = inspect_docx.inspect(src)["tables"]

    assert tables[0] == [["HELLO"]]


def test_inspect_surfaces_nested_table_text(tmp_path: Path) -> None:
    """Nested tables were silently dropped from inspect output. (#2154)"""
    from docx import Document

    inspect_docx = _inspect_docx_module()
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "outer"
    outer.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "inner"
    src = tmp_path / "nested.docx"
    doc.save(str(src))

    tables = inspect_docx.inspect(src)["tables"]

    assert "inner" in tables[0][0][0]


def test_inspect_does_not_mistake_a_field_code_for_a_tracked_change(tmp_path: Path) -> None:
    """``has_tracked_changes`` substring-matched the serialized body XML for
    ``"<w:ins"`` / ``"<w:del"``, which also matches ``<w:instrText>`` -- the
    field-instruction element behind an everyday PAGE/TOC/REF/hyperlink field
    -- and ``<w:insideH>``/``<w:insideV>`` -- a table's inside-border sides.
    Neither is a tracked change, and SKILL.md promises the flag only "when
    any w:ins or w:del element is found."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    inspect_docx = _inspect_docx_module()
    doc = Document()
    paragraph = doc.add_paragraph()

    def add_field(instr_text: str) -> None:
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldType"), "begin")
        paragraph.add_run()._r.append(begin)
        instr = OxmlElement("w:instrText")
        instr.text = instr_text
        paragraph.add_run()._r.append(instr)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldType"), "end")
        paragraph.add_run()._r.append(end)

    add_field("PAGE")

    table = doc.add_table(rows=1, cols=1)
    borders = OxmlElement("w:tblBorders")
    for side in ("insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        borders.append(edge)
    table._tbl.tblPr.append(borders)

    src = tmp_path / "field.docx"
    doc.save(str(src))

    result = inspect_docx.inspect(src)

    assert result["has_tracked_changes"] is False


def test_inspect_reports_a_real_tracked_change(tmp_path: Path) -> None:
    """The fix for the field-code false positive must not stop detecting an
    actual ``<w:ins>``/``<w:del>`` tracked change."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    inspect_docx = _inspect_docx_module()
    doc = Document()
    paragraph = doc.add_paragraph()
    ins = OxmlElement("w:ins")
    ins.set(qn("w:author"), "Reviewer")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "inserted"
    run.append(text)
    ins.append(run)
    paragraph._p.append(ins)

    src = tmp_path / "tracked.docx"
    doc.save(str(src))

    result = inspect_docx.inspect(src)

    assert result["has_tracked_changes"] is True


def test_inspect_rejects_non_docx_file(tmp_path: Path) -> None:
    """A corrupt file raised unhandled `PackageNotFoundError`; it must raise
    `ValueError` so the CLI can exit 2 with a clean message. (#2154)"""
    import pytest

    inspect_docx = _inspect_docx_module()
    bad = tmp_path / "invalid.docx"
    bad.write_text("invalid")

    with pytest.raises(ValueError, match="not a readable .docx file"):
        inspect_docx.inspect(bad)


def test_inspect_cli_rejects_non_docx_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    inspect_docx = _inspect_docx_module()
    bad = tmp_path / "invalid.docx"
    bad.write_text("invalid")

    monkeypatch.setattr(sys, "argv", ["inspect_docx.py", str(bad)])
    assert inspect_docx.main() == 2


def test_inspect_rejects_broken_package_xml(tmp_path: Path) -> None:
    """A zip that is not a valid OPC package (truncated `[Content_Types].xml`)
    leaked lxml's `XMLSyntaxError`; it must raise `ValueError` like every other
    unreadable file. (#2154)"""
    import zipfile

    inspect_docx = _inspect_docx_module()
    bad = tmp_path / "broken.docx"
    with zipfile.ZipFile(bad, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types")

    with pytest.raises(ValueError, match="not a readable .docx file"):
        inspect_docx.inspect(bad)


def test_inspect_docx_creates_parent_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "src.docx"
    create_docx.build({"body": [{"kind": "paragraph", "text": "x"}]}).save(str(src))

    out = tmp_path / "nested" / "dir" / "out.json"
    monkeypatch.setattr(sys, "argv", ["inspect_docx.py", str(src), "--out", str(out)])
    assert inspect_docx.main() == 0
    assert out.is_file()


def _edit_docx_module() -> object:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import edit_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return edit_docx


def _paragraph(runs: list[tuple[str, bool]]) -> object:
    """Build a one-paragraph document whose runs carry the given bold flags."""
    from docx import Document

    document = Document()
    paragraph = document.add_paragraph()
    for text, bold in runs:
        run = paragraph.add_run(text)
        if bold:
            run.bold = True
    return paragraph


def test_replace_text_keeps_the_formatting_of_untouched_runs() -> None:
    """A run the match never touched must survive byte-identical.

    Runs are where Word stores character formatting, so collapsing a paragraph
    into `runs[0]` gave every character run 0's formatting and left the rest as
    empty shells -- bold, italic, font and colour discarded for the whole
    paragraph even though one word changed.
    """
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("Hello ", False), ("world", True), (" and more", False)])

    assert edit_docx._replace_text_in_paragraph(paragraph, "Hello", "Hi") is True

    assert [(run.text, run.bold) for run in paragraph.runs] == [
        ("Hi ", None),
        ("world", True),
        (" and more", None),
    ]


def test_replace_text_across_a_run_boundary_keeps_the_second_run() -> None:
    """A `find` spanning runs still leaves the surrounding text in place.

    The replacement lands in the run owning the match's first character; the
    rest of the second run keeps its own text and formatting.
    """
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("Hel", False), ("lo world", True)])

    assert edit_docx._replace_text_in_paragraph(paragraph, "Hello", "Hi") is True

    assert [(run.text, run.bold) for run in paragraph.runs] == [("Hi", None), (" world", True)]


def test_replace_text_handles_several_matches_without_moving_text() -> None:
    """Each match is replaced where it starts, so runs keep their own share."""
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("aXa", False), ("Xa", True)])

    assert edit_docx._replace_text_in_paragraph(paragraph, "X", "-") is True

    assert [(run.text, run.bold) for run in paragraph.runs] == [("a-a", None), ("-a", True)]


@pytest.mark.parametrize(
    ("runs", "find", "replacement"),
    [
        ([("Hello ", False), ("world", True), (" and more", False)], "Hello", "Hi"),
        ([("Hel", False), ("lo world", True)], "Hello", "Hi"),
        ([("a", False), ("b", True), ("c", False)], "abc", "X"),
        ([("x{{N}}y", False)], "{{N}}", "Wei"),
        ([("aXa", False), ("Xa", True)], "X", "-"),
        ([("keep ", False), ("me", True)], "me", ""),
        ([("aa", False), ("aa", True)], "aa", "b"),
    ],
)
def test_replace_text_matches_str_replace_on_the_joined_text(
    runs: list[tuple[str, bool]], find: str, replacement: str
) -> None:
    """Whatever the run layout, the resulting text is plain `str.replace`.

    The old code already got the text right -- it was the run layout that was
    wrong -- so this pins the half that must not change while the fix moves
    characters back into their own runs.
    """
    edit_docx = _edit_docx_module()
    paragraph = _paragraph(runs)
    original = "".join(text for text, _ in runs)

    edit_docx._replace_text_in_paragraph(paragraph, find, replacement)

    assert "".join(run.text for run in paragraph.runs) == original.replace(find, replacement)


def test_replace_text_reports_false_when_the_needle_is_absent() -> None:
    """No match means no edit and no reported change."""
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("Hello ", False), ("world", True)])

    assert edit_docx._replace_text_in_paragraph(paragraph, "absent", "x") is False
    assert [(run.text, run.bold) for run in paragraph.runs] == [("Hello ", None), ("world", True)]


def test_replace_run_still_touches_only_its_own_run() -> None:
    """`replace_run` was never affected; keep it that way."""
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("Hello ", False), ("world", True)])

    edit_docx._replace_run(paragraph, 1, "there")

    assert [(run.text, run.bold) for run in paragraph.runs] == [("Hello ", None), ("there", True)]


def test_replace_text_reaches_table_cells(tmp_path: Path) -> None:
    """Placeholders in contracts and invoices usually live inside tables.

    `apply_ops` only walked `doc.paragraphs`, which python-docx limits to the
    body, so a `{{CLIENT}}` in a table cell was never replaced and the op
    reported zero applications.
    """
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("Agreement Header")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Client:"
    table.cell(0, 1).text = "{{CLIENT}}"

    applied = edit_docx.apply_ops(
        doc, [{"op": "replace_text", "find": "{{CLIENT}}", "with": "Acme Corp"}]
    )

    assert applied == 1
    assert table.cell(0, 1).text == "Acme Corp"
    assert table.cell(0, 0).text == "Client:"


def test_replace_text_counts_body_and_table_paragraphs_together() -> None:
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("Dear {{NAME}},")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Signed: {{NAME}}"

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{NAME}}", "with": "Wei"}])

    assert applied == 2
    assert doc.paragraphs[0].text == "Dear Wei,"
    assert doc.tables[0].cell(0, 0).text == "Signed: Wei"


def test_replace_text_reaches_nested_tables() -> None:
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Total: {{TOTAL}}"

    applied = edit_docx.apply_ops(
        doc, [{"op": "replace_text", "find": "{{TOTAL}}", "with": "42.00"}]
    )

    assert applied == 1
    assert inner.cell(0, 0).text == "Total: 42.00"


def test_replace_text_visits_a_merged_cell_once() -> None:
    """`row.cells` repeats a merged cell for every grid column it spans.

    Walking it once per column would apply the replacement again to text the
    first pass already rewrote; a replacement containing its own needle makes
    that visible.
    """
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 2))
    merged.text = "{{X}}"

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{X}}", "with": "{{X}}!"}])

    assert applied == 1
    assert merged.text == "{{X}}!"


def test_replace_text_survives_an_irregular_vertical_merge() -> None:
    """`row.cells` resolves a `vMerge=continue` cell against the row above and
    raises `ValueError` when no cell starts at that grid offset there -- a
    layout non-Word generators produce. Walking the `<w:tc>` elements directly
    never enters that path, so the whole edit (body included) still lands."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("Header {{X}}")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "Cell {{X}}"
    continue_marker = OxmlElement("w:vMerge")
    continue_marker.set(qn("w:val"), "continue")
    table.cell(1, 1)._tc.get_or_add_tcPr().append(continue_marker)

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{X}}", "with": "Y"}])

    assert applied == 2
    assert doc.paragraphs[0].text == "Header Y"
    assert table.cell(1, 0).text == "Cell Y"


def test_replace_text_reaches_section_headers_and_footers() -> None:
    """Letterheads and confidentiality banners live in headers and footers.

    `_iter_all_paragraphs` walked the body and its tables only, so a
    `{{ORG}}` in the header was never replaced and the op reported zero.
    """
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Header {{ORG}}"
    section.footer.paragraphs[0].text = "Footer {{CONFIDENTIAL}}"

    applied = edit_docx.apply_ops(
        doc,
        [
            {"op": "replace_text", "find": "{{ORG}}", "with": "Acme Corp"},
            {"op": "replace_text", "find": "{{CONFIDENTIAL}}", "with": "Public"},
        ],
    )

    assert applied == 2
    assert section.header.paragraphs[0].text == "Header Acme Corp"
    assert section.footer.paragraphs[0].text == "Footer Public"


def test_replace_text_reaches_first_and_even_page_headers_and_footers() -> None:
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = True
    section.first_page_header.paragraphs[0].text = "First {{X}}"
    section.first_page_footer.paragraphs[0].text = "First foot {{X}}"
    section.even_page_header.paragraphs[0].text = "Even {{X}}"
    section.even_page_footer.paragraphs[0].text = "Even foot {{X}}"

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{X}}", "with": "Y"}])

    assert applied == 4
    assert section.first_page_header.paragraphs[0].text == "First Y"
    assert section.first_page_footer.paragraphs[0].text == "First foot Y"
    assert section.even_page_header.paragraphs[0].text == "Even Y"
    assert section.even_page_footer.paragraphs[0].text == "Even foot Y"


def test_replace_text_reaches_tables_inside_headers() -> None:
    from docx import Document
    from docx.shared import Inches

    edit_docx = _edit_docx_module()
    doc = Document()
    header = doc.sections[0].header
    header.add_table(rows=1, cols=1, width=Inches(2)).cell(0, 0).text = "Ref {{REF}}"

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{REF}}", "with": "42"}])

    assert applied == 1
    assert header.tables[0].cell(0, 0).text == "Ref 42"


def test_replace_text_visits_a_header_shared_by_linked_sections_once() -> None:
    """A section linked to the previous one reuses that section's header part.

    Walking it again would replace nothing (the placeholder is already gone)
    but must not be counted twice, and must not create a header part on a
    section that has none: python-docx materialises one on first access.
    """
    from docx import Document
    from docx.enum.section import WD_SECTION

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Shared {{ORG}}"
    doc.add_section(WD_SECTION.NEW_PAGE)
    assert doc.sections[1].header.is_linked_to_previous is True

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{ORG}}", "with": "Acme"}])

    assert applied == 1
    assert doc.sections[0].header.paragraphs[0].text == "Shared Acme"
    assert doc.sections[1].header.is_linked_to_previous is True


def test_replace_text_does_not_materialise_absent_headers() -> None:
    """Touching `section.header.paragraphs` on a header-less section creates
    the header part; a body-only replacement must leave the package alone."""
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("Body {{X}}")

    edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{X}}", "with": "Y"}])

    section = doc.sections[0]
    assert section.header.is_linked_to_previous is True
    assert section.footer.is_linked_to_previous is True
    assert section.first_page_header.is_linked_to_previous is True
    assert section.even_page_footer.is_linked_to_previous is True


def test_apply_ops_skips_non_dict_ops() -> None:
    """Malformed op lists are ignored, as `edit_xlsx.apply_ops` already does."""
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("Hello {{NAME}}")

    applied = edit_docx.apply_ops(
        doc,
        [None, "replace_text", 3, {"op": "replace_text", "find": "{{NAME}}", "with": "Wei"}],
    )

    assert applied == 1
    assert doc.paragraphs[0].text == "Hello Wei"


def _two_paragraph_document() -> object:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello world")
    doc.add_paragraph("Second paragraph")
    return doc


def test_replace_run_reports_whether_it_wrote() -> None:
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("Hello ", False), ("world", True)])

    assert edit_docx._replace_run(paragraph, 1, "there") is True
    assert edit_docx._replace_run(paragraph, 2, "nope") is False
    assert edit_docx._replace_run(paragraph, -1, "nope") is False
    assert [run.text for run in paragraph.runs] == ["Hello ", "there"]


@pytest.mark.parametrize("run_idx", [1, 99, -1])
def test_apply_ops_does_not_count_an_out_of_bounds_run(run_idx: int) -> None:
    """A replace_run that wrote nothing must not be reported as applied (#1896)."""
    edit_docx = _edit_docx_module()
    doc = _two_paragraph_document()

    applied = edit_docx.apply_ops(
        doc, [{"op": "replace_run", "para": 0, "run": run_idx, "text": "New text"}]
    )

    assert applied == 0
    assert doc.paragraphs[0].text == "Hello world"


@pytest.mark.parametrize("para_idx", [2, 99, -1])
def test_apply_ops_does_not_count_an_out_of_bounds_paragraph(para_idx: int) -> None:
    """A negative index must not silently wrap to the end of the document."""
    edit_docx = _edit_docx_module()
    doc = _two_paragraph_document()

    applied = edit_docx.apply_ops(
        doc, [{"op": "replace_run", "para": para_idx, "run": 0, "text": "New text"}]
    )

    assert applied == 0
    assert [p.text for p in doc.paragraphs] == ["Hello world", "Second paragraph"]


def test_apply_ops_counts_only_the_replace_runs_that_wrote() -> None:
    edit_docx = _edit_docx_module()
    doc = _two_paragraph_document()

    applied = edit_docx.apply_ops(
        doc,
        [
            {"op": "replace_run", "para": 0, "run": 0, "text": "Hi world"},
            {"op": "replace_run", "para": 0, "run": 5, "text": "dropped"},
            {"op": "replace_run", "para": 1, "run": 0, "text": "Last paragraph"},
            {"op": "replace_run", "para": "x", "run": 0, "text": "dropped"},
        ],
    )

    assert applied == 2
    assert [p.text for p in doc.paragraphs] == ["Hi world", "Last paragraph"]


def _create_docx_module() -> object:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return create_docx


@pytest.mark.parametrize("spec", [None, [], "not-a-dict", 42, ("also", "not-a-dict")])
def test_build_returns_an_empty_document_for_a_non_dict_spec(spec: object) -> None:
    create_docx = _create_docx_module()

    doc = create_docx.build(spec)

    assert doc.paragraphs == []
    assert doc.tables == []


@pytest.mark.parametrize(
    "raw_body",
    [
        pytest.param("not-a-list", id="string"),
        pytest.param(None, id="null"),
        pytest.param([1, "x", None], id="list-of-non-dicts"),
    ],
)
def test_build_returns_an_empty_document_for_malformed_body(raw_body: object) -> None:
    create_docx = _create_docx_module()

    doc = create_docx.build({"body": raw_body})

    assert doc.paragraphs == []


def test_build_accepts_a_tuple_body_and_tuple_rows() -> None:
    """`body`/`rows` are consumed as any sequence, not specifically `list`."""
    create_docx = _create_docx_module()

    doc = create_docx.build({"body": ({"kind": "table", "rows": (["a", "b"], ["c", "d"])},)})

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["a", "b"]
    assert [c.text for c in table.rows[1].cells] == ["c", "d"]


def test_build_skips_a_table_whose_rows_are_all_empty() -> None:
    create_docx = _create_docx_module()

    doc = create_docx.build({"body": [{"kind": "table", "rows": [[], []]}]})

    assert doc.tables == []


def test_build_normalizes_scalar_table_rows_into_single_cell_rows() -> None:
    """A scalar row becomes one cell, not an iteration over its characters/digits."""
    create_docx = _create_docx_module()

    doc = create_docx.build({"body": [{"kind": "table", "rows": ["header", 42]}]})

    table = doc.tables[0]
    assert table.rows[0].cells[0].text == "header"
    assert table.rows[1].cells[0].text == "42"


@pytest.mark.parametrize("level", [-5, 50, None, "not-a-number"])
def test_build_clamps_an_invalid_heading_level(level: object) -> None:
    create_docx = _create_docx_module()

    # Must not raise -- python-docx itself only accepts levels 0-9.
    doc = create_docx.build({"body": [{"kind": "heading", "text": "Title", "level": level}]})

    assert doc.paragraphs[0].text == "Title"


def test_create_docx_cli_reports_invalid_json_with_exit_code_2(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    create_docx = _create_docx_module()

    bad_json = tmp_path / "bad.json"
    bad_json.write_text("{not valid json", encoding="utf-8")
    out = tmp_path / "out.docx"
    monkeypatch.setattr(sys, "argv", ["create_docx.py", str(bad_json), "--out", str(out)])

    assert create_docx.main() == 2
    assert "is not valid JSON" in capsys.readouterr().err
    assert not out.exists()


def test_create_docx_cli_reports_non_object_json_with_exit_code_2(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    create_docx = _create_docx_module()

    array_spec = tmp_path / "spec.json"
    array_spec.write_text('[{"kind": "paragraph", "text": "Hi"}]', encoding="utf-8")
    out = tmp_path / "out.docx"
    monkeypatch.setattr(sys, "argv", ["create_docx.py", str(array_spec), "--out", str(out)])

    assert create_docx.main() == 2
    assert 'must be a JSON object with a "body" array' in capsys.readouterr().err
    assert not out.exists()


# ── unusable ops and specs are reported, not swallowed (#2145, #2146) ────────


def _one_paragraph_docx(path: Path, text: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))


@pytest.mark.parametrize(
    ("body", "expected"),
    [
        ("not json at all", "is not valid JSON"),
        ('{"op": "replace_text", "find": "a", "with": "b"}', "must be a JSON array"),
        ('["replace_text"]', "op 0 must be an object"),
        ('[{"op": "replace-text", "find": "a", "with": "b"}]', "unknown kind 'replace-text'"),
        ('[{"find": "a", "with": "b"}]', "unknown kind None"),
        (
            '[{"op": "replace_text", "find": "a", "with": "b"}, {"op": "Replace_Text"}]',
            "op 1 has unknown kind 'Replace_Text'",
        ),
    ],
)
def test_edit_docx_reports_an_unusable_ops_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    body: str,
    expected: str,
) -> None:
    """Each of these either crashed with a traceback or was silently dropped,
    leaving an unmodified copy at --out reported as a successful edit."""
    edit_docx = _edit_docx_module()
    src = tmp_path / "in.docx"
    _one_paragraph_docx(src, "Hello ORIGINAL world")
    ops = tmp_path / "ops.json"
    ops.write_text(body, encoding="utf-8")
    out = tmp_path / "out.docx"
    monkeypatch.setattr(sys, "argv", ["edit_docx.py", str(src), str(ops), "--out", str(out)])

    assert edit_docx.main() == 2

    captured = capsys.readouterr()
    assert captured.err.startswith("error: ")
    assert "Traceback" not in captured.err
    assert expected in captured.err
    assert captured.out == ""
    assert not out.exists()


def test_edit_docx_reports_an_ops_file_that_is_not_utf8(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    edit_docx = _edit_docx_module()
    src = tmp_path / "in.docx"
    _one_paragraph_docx(src, "Hello")
    ops = tmp_path / "ops.json"
    ops.write_bytes(b"\xff\xfe[]")
    out = tmp_path / "out.docx"
    monkeypatch.setattr(sys, "argv", ["edit_docx.py", str(src), str(ops), "--out", str(out)])

    assert edit_docx.main() == 2
    assert "is not valid JSON" in capsys.readouterr().err
    assert not out.exists()


def test_edit_docx_leaves_an_existing_output_alone_when_the_ops_are_unusable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The destructive case: a typo'd op kind used to replace a document that
    already held real work with an unmodified copy of the input, and exit 0."""
    edit_docx = _edit_docx_module()
    src = tmp_path / "in.docx"
    _one_paragraph_docx(src, "Hello ORIGINAL world")
    out = tmp_path / "out.docx"
    _one_paragraph_docx(out, "PREVIOUSLY EDITED CONTENT")
    before = out.read_bytes()
    ops = tmp_path / "ops.json"
    ops.write_text('[{"op": "replace-text", "find": "a", "with": "b"}]', encoding="utf-8")
    monkeypatch.setattr(sys, "argv", ["edit_docx.py", str(src), str(ops), "--out", str(out)])

    assert edit_docx.main() == 2
    assert out.read_bytes() == before


def test_edit_docx_validates_the_ops_before_opening_the_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The ops file is checked first, so a bad ops file is named even when the
    input is not a document python-docx can open."""
    edit_docx = _edit_docx_module()
    src = tmp_path / "in.docx"
    src.write_bytes(b"this is not a docx")
    ops = tmp_path / "ops.json"
    ops.write_text('[{"op": "replace-text"}]', encoding="utf-8")
    monkeypatch.setattr(
        sys, "argv", ["edit_docx.py", str(src), str(ops), "--out", str(tmp_path / "o.docx")]
    )

    assert edit_docx.main() == 2


def test_edit_docx_load_ops_returns_the_ops_unchanged(tmp_path: Path) -> None:
    """Validation only inspects; the ops reach ``apply_ops`` exactly as written."""
    edit_docx = _edit_docx_module()
    ops = [{"op": "replace_text", "find": "a", "with": "b", "extra": 1}]
    path = tmp_path / "ops.json"
    path.write_text(json.dumps(ops), encoding="utf-8")

    assert edit_docx.load_ops(path) == ops


def test_edit_docx_op_kinds_match_what_apply_ops_handles(tmp_path: Path) -> None:
    """Every advertised kind is applied; the allow-list and the dispatcher
    cannot drift apart without this failing."""
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("alpha beta")
    ops = [
        {"op": "replace_text", "find": "alpha", "with": "ALPHA"},
        {"op": "replace_run", "para": 0, "run": 0, "text": "gamma"},
    ]
    assert {op["op"] for op in ops} == set(edit_docx.OP_KINDS)

    assert edit_docx.apply_ops(doc, ops) == 2


def test_edit_docx_an_empty_ops_array_is_still_accepted(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """``[]`` is a valid list of zero ops: the document is copied, applied is 0."""
    edit_docx = _edit_docx_module()
    src = tmp_path / "in.docx"
    _one_paragraph_docx(src, "Hello")
    ops = tmp_path / "ops.json"
    ops.write_text("[]", encoding="utf-8")
    out = tmp_path / "out.docx"
    monkeypatch.setattr(sys, "argv", ["edit_docx.py", str(src), str(ops), "--out", str(out)])

    assert edit_docx.main() == 0
    assert json.loads(capsys.readouterr().out) == {"applied": 0}
    assert out.is_file()


def test_edit_docx_still_applies_a_valid_ops_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """Validation must not disturb the working path."""
    from docx import Document

    edit_docx = _edit_docx_module()
    src = tmp_path / "in.docx"
    _one_paragraph_docx(src, "Hello ORIGINAL world")
    ops = tmp_path / "ops.json"
    ops.write_text(
        json.dumps([{"op": "replace_text", "find": "ORIGINAL", "with": "PATCHED"}]),
        encoding="utf-8",
    )
    out = tmp_path / "out.docx"
    monkeypatch.setattr(sys, "argv", ["edit_docx.py", str(src), str(ops), "--out", str(out)])

    assert edit_docx.main() == 0

    assert json.loads(capsys.readouterr().out) == {"applied": 1}
    assert [p.text for p in Document(str(out)).paragraphs] == ["Hello PATCHED world"]


@pytest.mark.parametrize(
    ("body", "expected"),
    [
        ('{"body": [{"kind": "paragrpah", "text": "Hi"}]}', "unknown kind 'paragrpah'"),
        ('{"body": [{"text": "Hi"}]}', "body entry 0 has unknown kind None"),
        ('{"body": ["paragraph"]}', "body entry 0 must be an object, got str"),
        (
            '{"body": [{"kind": "paragraph", "text": "ok"}, {"kind": "Heading", "text": "T"}]}',
            "body entry 1 has unknown kind 'Heading'",
        ),
    ],
)
def test_create_docx_reports_an_unrecognised_body_entry(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    body: str,
    expected: str,
) -> None:
    """A typo'd kind fell through ``build`` and produced an empty document at
    exit 0; now it is named by index, and nothing is written."""
    create_docx = _create_docx_module()
    spec = tmp_path / "spec.json"
    spec.write_text(body, encoding="utf-8")
    out = tmp_path / "new.docx"
    monkeypatch.setattr(sys, "argv", ["create_docx.py", str(spec), "--out", str(out)])

    assert create_docx.main() == 2

    captured = capsys.readouterr()
    assert captured.err.startswith("error: ")
    assert expected in captured.err
    if "unknown kind" in expected:
        assert "expected one of heading, paragraph, table, page_break" in captured.err
    assert captured.out == ""
    assert not out.exists()


def test_create_docx_accepts_an_empty_body(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """``{"body": []}`` is a valid empty document, not an unrecognised entry."""
    from docx import Document

    create_docx = _create_docx_module()
    spec = tmp_path / "spec.json"
    spec.write_text('{"body": []}', encoding="utf-8")
    out = tmp_path / "new.docx"
    monkeypatch.setattr(sys, "argv", ["create_docx.py", str(spec), "--out", str(out)])

    assert create_docx.main() == 0

    assert json.loads(capsys.readouterr().out) == {"entries": 0, "out": str(out)}
    assert Document(str(out)).paragraphs == []


def test_create_docx_body_kinds_match_what_build_renders() -> None:
    """Every advertised kind produces something; the allow-list and ``build``
    cannot drift apart without this failing."""
    create_docx = _create_docx_module()
    entries = [
        {"kind": "heading", "text": "T", "level": 1},
        {"kind": "paragraph", "text": "p"},
        {"kind": "table", "rows": [["a"]]},
        {"kind": "page_break"},
    ]
    assert [e["kind"] for e in entries] == list(create_docx.BODY_KINDS)

    create_docx.check_body_entries({"body": entries})
    doc = create_docx.build({"body": entries})

    assert [p.text for p in doc.paragraphs][:2] == ["T", "p"]
    assert len(doc.tables) == 1


def test_create_docx_still_builds_a_valid_spec_and_now_reports_it(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    """`create_docx` printed nothing at all on success, so a caller had no
    signal beyond the exit code."""
    from docx import Document

    create_docx = _create_docx_module()
    spec = tmp_path / "spec.json"
    spec.write_text(
        json.dumps(
            {"body": [{"kind": "heading", "text": "T"}, {"kind": "paragraph", "text": "Hi"}]}
        ),
        encoding="utf-8",
    )
    out = tmp_path / "new.docx"
    monkeypatch.setattr(sys, "argv", ["create_docx.py", str(spec), "--out", str(out)])

    assert create_docx.main() == 0

    assert json.loads(capsys.readouterr().out) == {"entries": 2, "out": str(out)}
    assert [p.text for p in Document(str(out)).paragraphs] == ["T", "Hi"]


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_V = "urn:schemas-microsoft-com:vml"
_WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


def _txbx_content(inner: str) -> str:
    return f"<w:txbxContent>{inner}</w:txbxContent>"


def _paragraph_xml(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def _add_vml_text_box(para: object, inner: str) -> None:
    """Attach a legacy ``<w:pict>`` text box holding *inner* to *para*."""
    from lxml import etree

    run = etree.fromstring(
        f'<w:r xmlns:w="{_W}" xmlns:v="{_V}"><w:pict><v:shape><v:textbox>'
        f"{_txbx_content(inner)}"
        "</v:textbox></v:shape></w:pict></w:r>"
    )
    para._p.append(run)


def _add_drawingml_text_box(para: object, inner: str) -> None:
    """Attach a modern ``<w:drawing>`` shape text box holding *inner* to *para*."""
    from lxml import etree

    run = etree.fromstring(
        f'<w:r xmlns:w="{_W}" xmlns:wps="{_WPS}"><w:drawing><wps:wsp><wps:txbx>'
        f"{_txbx_content(inner)}"
        "</wps:txbx></wps:wsp></w:drawing></w:r>"
    )
    para._p.append(run)


def _document_xml(doc: object, tmp_path: Path) -> str:
    import zipfile

    out = tmp_path / "saved.docx"
    doc.save(str(out))
    return zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")


def test_replace_text_reaches_a_vml_text_box(tmp_path: Path) -> None:
    """A text box's words belong to no paragraph any walk reached.

    Word parks text-box content in a ``<w:txbxContent>`` nested inside a run, so
    python-docx reports the host paragraph as empty. A `replace_text` that
    rewrote the body and left the box alone still reported the body hit, which
    reads as a completed redaction.
    """
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("Prepared for ACME-CODENAME under NDA.")
    _add_vml_text_box(doc.add_paragraph(), _paragraph_xml("Confidential: ACME-CODENAME"))

    applied = edit_docx.apply_ops(
        doc, [{"op": "replace_text", "find": "ACME-CODENAME", "with": "Project Harbor"}]
    )

    assert applied == 2
    xml = _document_xml(doc, tmp_path)
    assert "ACME-CODENAME" not in xml
    assert xml.count("Project Harbor") == 2


def test_replace_text_reaches_a_drawingml_text_box(tmp_path: Path) -> None:
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    _add_drawingml_text_box(doc.add_paragraph(), _paragraph_xml("Callout: {{NAME}}"))

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{NAME}}", "with": "Wei"}])

    assert applied == 1
    xml = _document_xml(doc, tmp_path)
    assert "{{NAME}}" not in xml
    assert "Callout: Wei" in xml


def test_replace_text_reaches_a_text_box_in_a_header(tmp_path: Path) -> None:
    """Letterhead banners are text boxes inside a header — two misses at once."""
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    header = doc.sections[0].header
    header.is_linked_to_previous = False
    _add_vml_text_box(header.paragraphs[0], _paragraph_xml("DRAFT — {{STATUS}}"))

    applied = edit_docx.apply_ops(
        doc, [{"op": "replace_text", "find": "{{STATUS}}", "with": "Final"}]
    )

    assert applied == 1
    assert "DRAFT — Final" in header.part.blob.decode("utf-8")


def test_replace_text_reaches_a_table_inside_a_text_box(tmp_path: Path) -> None:
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    _add_vml_text_box(
        doc.add_paragraph(),
        f"<w:tbl><w:tr><w:tc>{_paragraph_xml('Total: {{TOTAL}}')}</w:tc></w:tr></w:tbl>",
    )

    applied = edit_docx.apply_ops(
        doc, [{"op": "replace_text", "find": "{{TOTAL}}", "with": "42.00"}]
    )

    assert applied == 1
    assert "Total: 42.00" in _document_xml(doc, tmp_path)


def test_replace_text_reaches_a_text_box_inside_a_text_box(tmp_path: Path) -> None:
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    inner_box = (
        f'<w:p><w:r xmlns:v="{_V}"><w:pict><v:shape><v:textbox>'
        f"{_txbx_content(_paragraph_xml('deep {{X}}'))}"
        "</v:textbox></v:shape></w:pict></w:r></w:p>"
    )
    _add_vml_text_box(doc.add_paragraph(), inner_box)

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{X}}", "with": "found"}])

    assert applied == 1
    assert "deep found" in _document_xml(doc, tmp_path)


def test_replace_text_visits_a_text_box_paragraph_once(tmp_path: Path) -> None:
    """A replacement containing its own needle exposes a paragraph visited twice."""
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    _add_vml_text_box(doc.add_paragraph(), _paragraph_xml("{{X}}"))

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{X}}", "with": "{{X}}!"}])

    assert applied == 1
    assert "{{X}}!" in _document_xml(doc, tmp_path)
    assert "{{X}}!!" not in _document_xml(doc, tmp_path)


def test_replace_run_still_indexes_only_body_paragraphs() -> None:
    """Guard: text boxes must not shift the index space `replace_run` addresses."""
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("first")
    _add_vml_text_box(doc.add_paragraph(), _paragraph_xml("boxed"))
    doc.add_paragraph("third")

    applied = edit_docx.apply_ops(doc, [{"op": "replace_run", "para": 2, "run": 0, "text": "3rd"}])

    assert applied == 1
    assert doc.paragraphs[2].text == "3rd"

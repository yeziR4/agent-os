"""Create a `.xlsx` workbook from a JSON spec.

Spec:
    {
      "sheets": [
        {
          "name": "Sales",
          "rows": [["A", "B"], [1, "=B1*2"]],
          "merged": [{"range": "A1:B1"}],
          "freeze": "A2"
        }
      ]
    }
Entries in "merged" may also be range strings, e.g. "A1:B1", as returned by
inspect_xlsx. Strings and {"range": ...} objects can be mixed in the same list.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook


class SpecError(ValueError):
    """A spec that cannot be used. Reported as ``error:`` / exit 2, never as a
    traceback: the caller passed bad input, the script did not break."""


def check_sheets_shape(spec: dict[str, Any]) -> None:
    """Raise :class:`SpecError` for a "sheets" shape ``build`` would silently drop.

    A "sheets" value that is not a list, or a list entry that is not an
    object, used to fall through ``build``'s ``isinstance`` filters
    unmentioned: ``{"sheets": [1, 2, 3]}`` and ``{"sheets": "Sales"}`` both
    produced a valid, empty, single-default-sheet workbook reported as a
    success, indistinguishable from an intentional blank spec. A missing
    "sheets" key is left alone -- an omitted key is a deliberate blank
    workbook, not a malformed one.
    """
    if "sheets" not in spec:
        return
    raw_sheets = spec["sheets"]
    if not isinstance(raw_sheets, (list, tuple)):
        raise SpecError(f'"sheets" must be an array, got {type(raw_sheets).__name__}')
    for index, item in enumerate(raw_sheets):
        if not isinstance(item, dict):
            raise SpecError(f"sheet entry {index} must be an object, got {type(item).__name__}")


def _coerce(value: Any) -> Any:
    if isinstance(value, str) and len(value) >= 19 and value[10] == "T":
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return value
    return value


def build(spec: Any) -> Workbook:
    wb = Workbook()
    if not isinstance(spec, dict):
        return wb

    default_sheet = wb.active
    raw_sheets = spec.get("sheets")
    if not isinstance(raw_sheets, (list, tuple)):
        return wb
    sheets = [s for s in raw_sheets if isinstance(s, dict)]
    if not sheets:
        return wb

    for idx, sheet_spec in enumerate(sheets):
        if idx == 0:
            ws = default_sheet
            ws.title = str(sheet_spec.get("name") or "Sheet1")
        else:
            ws = wb.create_sheet(title=str(sheet_spec.get("name") or f"Sheet{idx + 1}"))

        raw_rows = sheet_spec.get("rows")
        if isinstance(raw_rows, (list, tuple)):
            for row in raw_rows:
                # A scalar row (str/int/None/...) becomes a single-cell row
                # rather than being iterated -- a bare string would otherwise
                # be split into one cell per character, and a non-iterable
                # scalar like an int would raise TypeError outright.
                values = row if isinstance(row, (list, tuple)) else [row]
                ws.append([_coerce(v) for v in values])

        raw_merged = sheet_spec.get("merged")
        if isinstance(raw_merged, (list, tuple)):
            for merged in raw_merged:
                if isinstance(merged, str):
                    ws.merge_cells(merged)
                elif (
                    isinstance(merged, dict)
                    and "range" in merged
                    and isinstance(merged["range"], str)
                ):
                    ws.merge_cells(merged["range"])

        freeze = sheet_spec.get("freeze")
        if isinstance(freeze, str) and freeze:
            ws.freeze_panes = freeze

    return wb


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create a .xlsx from a JSON spec.")
    parser.add_argument("spec", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    if not args.spec.is_file():
        print(f"error: spec {args.spec} not found", file=sys.stderr)
        return 2
    try:
        raw_spec = args.spec.read_text(encoding="utf-8")
        spec = json.loads(raw_spec)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        print(f"error: invalid JSON spec: {exc}", file=sys.stderr)
        return 2
    if not isinstance(spec, dict):
        print("error: JSON spec must be an object", file=sys.stderr)
        return 2
    try:
        check_sheets_shape(spec)
    except SpecError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    wb = build(spec)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(args.out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
